# -- coding: utf-8 --

from odoo import models, fields, api, _
import base64, csv
import io
from io import StringIO
from openpyxl import load_workbook
from odoo.exceptions import ValidationError

class MecaluxQuantFormat(models.TransientModel):
    _name = 'mecalux.quant.format'

    mecalux_file = fields.Binary('Xlsx File', attachment=True)
    mecalux_file_name = fields.Char(string="Filename")
    stock_location_id = fields.Many2one('stock.location', string="Stock Location", domain="[('usage', '=', 'internal'), '|', ('company_id', '=', False), ('company_id', '=', company_id)]", required=True)
    company_id = fields.Many2one(
        'res.company',
        string="Company",
        default=lambda self: self.env.company,
        required=True
    )

    def import_mecalux_quants(self):
        for rec in self:
            if not rec.mecalux_file:
                raise ValidationError("Please upload a file.")

            file_data = base64.b64decode(rec.mecalux_file)
            excel_file = io.BytesIO(file_data)

            try:
                wb = load_workbook(excel_file, data_only=True)
            except Exception as e:
                raise ValidationError(f"Unable to read the Excel file: {str(e)}")

            sheet = wb.active

            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                raise ValidationError("File is empty.")

            headers = rows[0]
            if not headers:
                raise ValidationError("No header found in imported file.")

            location_id = self.stock_location_id
            # location_id = self.env['stock.location'].search([('usage', '=', 'internal'), ('name', '=', 'MANUFAST'), ('location_id.name', '=', 'WH')], limit=1)
            if not location_id:
                raise ValidationError("No internal stock location found.")

            data_list = []
            non_existing_products = []

            for row_num, row in enumerate(rows[1:], start=2):  # Start from second row
                if not row or all(cell is None for cell in row):
                    continue  # skip empty rows

                product_ref = str(row[0]).strip() if row[0] else ''
                counted_qty = str(row[1]).strip()
                user_name = str(row[3]).strip() if len(row) > 3 and row[3] else ''

                if not product_ref:
                    raise ValidationError(f"Row {row_num}: Required value 'Product' is missing.")
                if not counted_qty:
                    raise ValidationError(f"Row {row_num}: Required value 'Quantity' is missing.")
                if not user_name:
                    raise ValidationError(f"Row {row_num}: Required value 'User' is missing.")

                product_id = self.env['product.product'].search([('default_code', '=', product_ref)], limit=1)
                user_id = self.env['res.users'].search([('name', '=', user_name)], limit=1)

                if product_id:
                    self.env['stock.quant'].search([('product_id', '=', product_id.id)]).action_set_inventory_quantity_zero()
                    data_list.append({
                        'product_id': product_id.id,
                        'quantity': float(counted_qty),
                        'user_id': user_id.id or self.env.user.id,
                        'location_id': location_id.id,
                    })
                else:
                    non_existing_products.append(product_ref)
                    
            if non_existing_products:
                raise ValidationError(f"The following product references do not exist: {', '.join(non_existing_products)}")

            if data_list:
                self.env['stock.quant'].create(data_list)
                return {
                    'type': 'ir.actions.client',
                    'tag': 'display_notification',
                    'params': {
                        'title': 'Mecalux Stock Quants Import',
                        'message': _('All stock quants are imported successfully.'),
                        'sticky': False,
                        'type': 'success',
                        'next': {'type': 'ir.actions.act_window_close'},
                    }
                }
            else:
                return {
                    'type': 'ir.actions.client',
                    'tag': 'display_notification',
                    'params': {
                        'title': 'Mecalux Stock Quants Import',
                        'message': _('No records to import!'),
                        'sticky': False,
                        'type': 'danger',
                        'next': {'type': 'ir.actions.act_window_close'},
                    }
                }