import os
from io import BytesIO

from odoo import http
from odoo.exceptions import AccessError
from odoo.http import request
from odoo.modules.module import get_module_resource

import openpyxl


class OpmsImportTemplateController(http.Controller):
    _SOURCE_FILENAME = "KPI 2025-26_Remapped-UPDATED.xlsx"
    _DOWNLOAD_FILENAME = "KPI 2025-26_Remapped-UPDATED.xlsx"
    _UNIT_HEADER = "Measurement Unit"
    _UNIT_HEADER_ALIASES = {"Measurement Unit", "Measurement_Unit", "Unit", "UOM"}

    def _build_template_payload_with_unit_column(self, payload):
        """Ensure the downloaded template contains a dedicated Measurement Unit column.

        This keeps old source templates usable while providing the new optional import field.
        """
        workbook = openpyxl.load_workbook(BytesIO(payload))
        sheet = workbook[workbook.sheetnames[0]]

        headers = []
        for col in range(1, sheet.max_column + 1):
            value = sheet.cell(row=1, column=col).value
            headers.append((str(value).strip() if value else ""))

        if any(header in self._UNIT_HEADER_ALIASES for header in headers):
            return payload

        insert_after_aliases = {
            "Comparison Direction",
            "Comparison_Direction",
            "Aggregation Rule",
            "Aggregation_Rule",
            "Measurement Type",
            "Measurement_Type",
        }
        insert_after_idx = -1
        for idx, header in enumerate(headers):
            if header in insert_after_aliases:
                insert_after_idx = max(insert_after_idx, idx)

        if insert_after_idx >= 0:
            insert_col = insert_after_idx + 2
            sheet.insert_cols(insert_col, amount=1)
            sheet.cell(row=1, column=insert_col).value = self._UNIT_HEADER
        else:
            sheet.cell(row=1, column=sheet.max_column + 1).value = self._UNIT_HEADER

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @http.route(["/opms/import/template/download"], type="http", auth="user")
    def opms_download_import_template(self, **kwargs):
        user = request.env.user
        if not user.has_group("base.group_user"):
            raise AccessError("Only internal users can download OPMS import templates.")

        file_path = get_module_resource("opms_ecdhs", "data", self._SOURCE_FILENAME)
        if not file_path or not os.path.exists(file_path):
            return request.not_found()

        with open(file_path, "rb") as template_file:
            payload = template_file.read()

        try:
            payload = self._build_template_payload_with_unit_column(payload)
        except Exception:
            # Never block downloads if workbook transformation fails.
            pass

        return request.make_response(
            payload,
            headers=[
                (
                    "Content-Type",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
                (
                    "Content-Disposition",
                    f'attachment; filename="{self._DOWNLOAD_FILENAME}"',
                ),
                ("Cache-Control", "no-store"),
                ("X-Content-Type-Options", "nosniff"),
            ],
        )