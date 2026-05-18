from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Inches


BASE = Path('/Users/benjaminmaimba/Documents/GitHub/jhbproperty/tender_management/docs')
DIAGRAMS = BASE / 'diagrams'
OUT = BASE / 'Tender_Management_User_and_Technical_Guide_BPMN_Graphics_Embedded.docx'


def build_doc() -> None:
    doc = Document()
    doc.add_heading('Tender Management Module', level=0)
    doc.add_paragraph('BPMN Swimlane Process Diagrams (Graphical Edition)')
    doc.add_paragraph(f'Date: {date.today().isoformat()}')
    doc.add_paragraph(
        'These diagrams show end-to-end process flows across users, frontend routes, '
        'backend controllers, and model workflow transitions.'
    )

    flows = [
        ('Diagram A: Tender Authoring to Publication', 'flow_a_tender_publication.png'),
        ('Diagram B: Public Tender Discovery and Detail Access', 'flow_b_discovery_detail.png'),
        ('Diagram C: Countdown Timer (Frontend to Backend JSON)', 'flow_c_timer.png'),
        ('Diagram D: Bid Submission End to End', 'flow_d_bid_submission.png'),
        ('Diagram E: Secure Document Download', 'flow_e_downloads.png'),
        ('Diagram F: Vendor Registration', 'flow_f_vendor_registration.png'),
        ('Diagram G: Portal Bid Tracking and Detail Access', 'flow_g_portal_tracking.png'),
        ('Diagram H: Close, Adjudication, and Award', 'flow_h_close_award.png'),
    ]

    for title, image_name in flows:
        image_path = DIAGRAMS / image_name
        doc.add_heading(title, level=2)
        if image_path.exists():
            doc.add_picture(str(image_path), width=Inches(6.8))
            doc.add_paragraph(f'Figure: {title}')
        else:
            doc.add_paragraph(f'Missing image: {image_name}')

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build_doc()
