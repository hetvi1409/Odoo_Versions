from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Inches


BASE = Path('/Users/benjaminmaimba/Documents/GitHub/jhbproperty/tender_management/docs')
DIAGRAMS = BASE / 'diagrams'

USER_OUT = BASE / 'Tender_Management_User_Guide_BPMN.docx'
TECH_OUT = BASE / 'Tender_Management_Technical_Guide_BPMN.docx'
USER_HTML_OUT = BASE / 'Tender_Management_User_Guide_BPMN.html'
TECH_HTML_OUT = BASE / 'Tender_Management_Technical_Guide_BPMN.html'

FLOWS = [
    ('Diagram A: Tender Authoring to Publication', 'flow_a_tender_publication.png'),
    ('Diagram B: Public Tender Discovery and Detail Access', 'flow_b_discovery_detail.png'),
    ('Diagram C: Countdown Timer', 'flow_c_timer.png'),
    ('Diagram D: Bid Submission End to End', 'flow_d_bid_submission.png'),
    ('Diagram E: Secure Document Download', 'flow_e_downloads.png'),
    ('Diagram F: Vendor Registration', 'flow_f_vendor_registration.png'),
    ('Diagram G: Portal Bid Tracking', 'flow_g_portal_tracking.png'),
    ('Diagram H: Close, Adjudication, and Award', 'flow_h_close_award.png'),
]


def add_flow_images(doc: Document, title: str) -> None:
    doc.add_heading(title, level=1)
    doc.add_paragraph('The following BPMN-style process diagrams are included for visual flow understanding.')
    for flow_title, image_name in FLOWS:
        image_path = DIAGRAMS / image_name
        doc.add_heading(flow_title, level=2)
        if image_path.exists():
            doc.add_picture(str(image_path), width=Inches(6.8))
            doc.add_paragraph(f'Figure: {flow_title}')
        else:
            doc.add_paragraph(f'Missing image: {image_name}')


def _html_header(title: str) -> str:
    return f"""<!DOCTYPE html>
<html>
<head>
  <meta charset=\"utf-8\" />
  <title>{title}</title>
  <style>
    body {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; line-height: 1.45; color: #111; max-width: 1100px; margin: 24px auto; padding: 0 16px; }}
    h1, h2, h3 {{ color: #0b3558; }}
    h1 {{ font-size: 28px; margin-bottom: 4px; }}
    h2 {{ font-size: 22px; margin-top: 28px; }}
    h3 {{ font-size: 18px; margin-top: 20px; }}
    .meta {{ color: #444; margin-bottom: 18px; }}
    .card {{ background: #f7fbff; border: 1px solid #dce8f5; border-radius: 8px; padding: 12px; margin: 10px 0; }}
    .flow-img {{ width: 100%; max-width: 1050px; border: 1px solid #cfcfcf; border-radius: 4px; padding: 4px; background: #fff; }}
    code {{ font-family: Consolas, 'Courier New', monospace; font-size: 10pt; background: #f2f2f2; padding: 1px 4px; border-radius: 3px; }}
    ul {{ margin-top: 6px; }}
  </style>
</head>
<body>
"""


def _html_bpmn_section() -> str:
    parts = [
        "<h2>Graphical BPMN Process Flows</h2>",
        "<p>The following BPMN-style process diagrams are included for visual flow understanding.</p>",
    ]
    for flow_title, image_name in FLOWS:
        parts.append(f"<h3>{flow_title}</h3>")
        image_path = DIAGRAMS / image_name
        if image_path.exists():
            parts.append(
                f"<p><img class=\"flow-img\" src=\"diagrams/{image_name}\" alt=\"{flow_title}\" /></p>"
            )
            parts.append(f"<p><em>Figure: {flow_title}</em></p>")
        else:
            parts.append(f"<p><strong>Missing image:</strong> {image_name}</p>")
    return "\n".join(parts)


def build_user_html() -> None:
    today = date.today().isoformat()
    html = [
        _html_header('Tender Management User Guide (BPMN)'),
        "<h1>Tender Management User Guide</h1>",
        "<p class=\"meta\"><strong>Audience:</strong> Procurement and vendor users (non-technical)<br/>"
        f"<strong>Date:</strong> {today}</p>",
        "<h2>1. What This System Is For</h2>",
        "<div class=\"card\">This system helps your organization publish tenders, receive bid submissions, and track bid outcomes in one place. It also allows vendors to register and submit bids online.</div>",
        "<h2>2. Who Uses This System</h2>",
        "<h3>Internal Procurement Team</h3>",
        "<ul><li>Create and publish tenders</li><li>Review submissions and move tenders through approval and award stages</li></ul>",
        "<h3>Vendors</h3>",
        "<ul><li>Register as vendors</li><li>View open tenders and submit bid documents</li><li>Track submitted bids in the portal</li></ul>",
        "<h2>3. User Journey - Internal Team</h2>",
        "<ul><li><strong>Step 1:</strong> Create Tender</li><li><strong>Step 2:</strong> Upload RFQ Documents</li><li><strong>Step 3:</strong> Move statuses: Submit -> Verify -> Approve -> Publish</li><li><strong>Step 4:</strong> Enable bidding when ready</li><li><strong>Step 5:</strong> Manage adjudication and award after closure</li></ul>",
        "<h2>4. User Journey - Vendor</h2>",
        "<ul><li><strong>Step 1:</strong> Register as vendor</li><li><strong>Step 2:</strong> Find a tender from RFQ/RFP pages</li><li><strong>Step 3:</strong> Review tender details and documents</li><li><strong>Step 4:</strong> Upload bid documents and submit</li><li><strong>Step 5:</strong> Track your submitted bid in portal</li></ul>",
        "<h2>5. Business Rules Users Must Know</h2>",
        "<ul><li>RFQ documents are mandatory before key status changes</li><li>Bid submission works only when bidding is enabled and tender is still open</li><li>Submission closes based on closing date/time and status</li><li>Bid document upload is mandatory</li></ul>",
        "<h2>6. Common Errors and What To Do</h2>",
        "<ul><li><strong>Upload is mandatory:</strong> Attach at least one supported document and retry</li><li><strong>Tender not available:</strong> Tender may be closed, unpublished, or bidding disabled</li><li><strong>Not authorized:</strong> Use the correct account for that bid/tender</li><li><strong>Timer N/A:</strong> Refresh page, then contact support if persistent</li></ul>",
        _html_bpmn_section(),
        "<h2>7. Support Handover</h2>",
        "<p>When reporting issues, include tender reference number, user account email, and a screenshot.</p>",
        "</body></html>",
    ]
    USER_HTML_OUT.write_text("\n".join(html), encoding='utf-8')


def build_technical_html() -> None:
    today = date.today().isoformat()
    html = [
        _html_header('Tender Management Technical Guide (BPMN)'),
        "<h1>Tender Management Technical Guide</h1>",
        "<p class=\"meta\"><strong>Audience:</strong> Developers, technical support, QA<br/>"
        f"<strong>Date:</strong> {today}</p>",
        "<h2>1. Module Overview</h2>",
        "<div class=\"card\"><strong>Module:</strong> <code>tender_management</code> for Odoo 16. Covers tender publication, bid submission, portal tracking, workflow control, downloads, and website integration.</div>",
        "<h2>2. Architecture</h2>",
        "<ul><li><strong>Controllers:</strong> <code>controller/bid_submission.py</code>, <code>controller/vendor_registration.py</code>, <code>controller/tender_bid_portal.py</code></li><li><strong>Models:</strong> <code>models/tender_tender.py</code>, <code>models/tender_bid.py</code>, <code>models/tender_document.py</code></li><li><strong>Views:</strong> backend forms and website templates under <code>views/</code></li><li><strong>Assets:</strong> <code>static/src/js/website_tender.js</code>, <code>static/src/css/tender_information.css</code></li></ul>",
        "<h2>3. Security and Validation Controls</h2>",
        "<ul><li>Bid submit route uses auth user + POST + CSRF</li><li>Vendor registration submit uses POST + CSRF + rate limit</li><li>Portal routes use owner/internal authorization checks</li><li>Download routes validate entitlement before streaming</li><li>Timer route validates HMAC token against <code>database.secret</code></li><li>Upload validation enforces extension and size limits</li></ul>",
        "<h2>4. Data Integrity and Locking</h2>",
        "<ul><li><code>tender.tender.write</code> blocks non-draft edits except controlled fields</li><li><code>tender.bid.write</code> blocks non-draft edits except controlled fields</li><li><code>tender.document</code> create/write/unlink blocked when linked tender is non-draft</li><li><code>fields_get</code> readonly metadata lock applied to non-draft records</li></ul>",
        "<h2>5. Configuration Keys</h2>",
        "<ul><li><code>tender_management.folder_id</code></li><li><code>tender_management.rate_limit.submit_bid.max_requests</code></li><li><code>tender_management.rate_limit.submit_bid.window_seconds</code></li><li><code>tender_management.rate_limit.vendor_registration.max_requests</code></li><li><code>tender_management.rate_limit.vendor_registration.window_seconds</code></li></ul>",
        "<h2>6. Troubleshooting</h2>",
        "<ul><li>Timer issues: verify <code>tender_timer_token</code> in page and JSON payload</li><li>Unauthorized download: validate publication state or bid ownership</li><li>Submission blocked: inspect tender status, <code>allow_bidding</code>, auth context, and upload rules</li><li>Rate-limit behavior: verify deployment topology and per-worker limits</li></ul>",
        _html_bpmn_section(),
        "<h2>7. File Reference Summary</h2>",
        "<p>Implementation resides in <code>controller/</code>, <code>models/</code>, <code>views/</code>, <code>security/</code>, and <code>static/src/</code>.</p>",
        "</body></html>",
    ]
    TECH_HTML_OUT.write_text("\n".join(html), encoding='utf-8')


def build_user_guide() -> None:
    doc = Document()
    doc.add_heading('Tender Management User Guide', level=0)
    doc.add_paragraph('Audience: Procurement and vendor users (non-technical)')
    doc.add_paragraph(f'Date: {date.today().isoformat()}')

    doc.add_heading('1. What This System Is For', level=1)
    doc.add_paragraph(
        'This system helps your organization publish tenders, receive bid submissions, '
        'and track bid outcomes in one place. It also allows vendors to register and submit bids online.'
    )

    doc.add_heading('2. Who Uses This System', level=1)
    doc.add_paragraph('Internal Procurement Team')
    doc.add_paragraph('- Create and publish tenders')
    doc.add_paragraph('- Review submissions and move tenders through approval and award stages')
    doc.add_paragraph('Vendors')
    doc.add_paragraph('- Register as vendors')
    doc.add_paragraph('- View open tenders and submit bid documents')
    doc.add_paragraph('- Track submitted bids in the portal')

    doc.add_heading('3. User Journey - Internal Team', level=1)
    doc.add_paragraph('Step 1: Create Tender')
    doc.add_paragraph('Enter tender details such as name, category, opening date, closing date, and type (RFQ/RFP).')
    doc.add_paragraph('Step 2: Upload RFQ Documents')
    doc.add_paragraph('Upload required RFQ documents before moving to the next status.')
    doc.add_paragraph('Step 3: Move Through Statuses')
    doc.add_paragraph('Submit -> Verify -> Approve -> Publish.')
    doc.add_paragraph('Step 4: Enable Vendor Bidding')
    doc.add_paragraph('Enable bidding for the tender when you want vendors to start submitting bids.')
    doc.add_paragraph('Step 5: Monitor Lifecycle')
    doc.add_paragraph('After closing date, manage adjudication and award statuses.')

    doc.add_heading('4. User Journey - Vendor', level=1)
    doc.add_paragraph('Step 1: Register as Vendor')
    doc.add_paragraph('Complete the vendor registration form and submit your details.')
    doc.add_paragraph('Step 2: Find a Tender')
    doc.add_paragraph('Browse RFQ/RFP listings and open a tender detail page.')
    doc.add_paragraph('Step 3: Review Details and Documents')
    doc.add_paragraph('Read tender information, download RFQ and related documents, and check countdown timer.')
    doc.add_paragraph('Step 4: Submit Bid')
    doc.add_paragraph('Upload required bid documents and submit your bid.')
    doc.add_paragraph('Step 5: Track Your Bid')
    doc.add_paragraph('Use the portal to view your submitted bids and status updates.')

    doc.add_heading('5. Key Screens and What They Do', level=1)
    doc.add_paragraph('Tender Listings')
    doc.add_paragraph('- Shows available RFQs and RFPs with search and paging.')
    doc.add_paragraph('Tender Detail')
    doc.add_paragraph('- Shows full tender info, briefing details, downloadable documents, and submission actions.')
    doc.add_paragraph('Bid Submission Page')
    doc.add_paragraph('- Lets logged-in vendors upload documents and submit bids.')
    doc.add_paragraph('Bid Status and Portal Pages')
    doc.add_paragraph('- Shows tender outcomes and your own submitted bid records.')

    doc.add_heading('6. Business Rules Users Must Know', level=1)
    doc.add_paragraph('- RFQ documents are mandatory before key status changes.')
    doc.add_paragraph('- Vendors can submit bids only when bidding is enabled and tender is still open.')
    doc.add_paragraph('- Submission closes automatically based on closing date/time and status.')
    doc.add_paragraph('- Uploaded bid documents are required for submission.')

    doc.add_heading('7. Common Errors and What To Do', level=1)
    doc.add_paragraph('"Upload Documents is mandatory before submission"')
    doc.add_paragraph('- Attach one or more supported documents before submitting.')
    doc.add_paragraph('"This tender is not available for bid submission"')
    doc.add_paragraph('- Tender may be closed, not published, or bidding is disabled.')
    doc.add_paragraph('"You are not authorized to download this document"')
    doc.add_paragraph('- Ensure you are using the correct account and have access to the tender or bid.')
    doc.add_paragraph('Timer shows N/A')
    doc.add_paragraph('- Refresh the page; if issue continues, report to support team.')

    doc.add_heading('8. User Checklist Before Go-Live', level=1)
    doc.add_paragraph('- All tender templates and documents prepared')
    doc.add_paragraph('- Approval path understood by internal users')
    doc.add_paragraph('- Vendor communication sent with portal instructions')
    doc.add_paragraph('- Support contacts shared for registration and submission issues')

    add_flow_images(doc, '9. Graphical BPMN Process Flows')

    doc.add_heading('10. Support Handover Notes', level=1)
    doc.add_paragraph('If users experience access, download, or submission issues, log the tender reference and screenshot, then escalate to the technical support team.')

    doc.save(USER_OUT)


def build_technical_guide() -> None:
    doc = Document()
    doc.add_heading('Tender Management Technical Guide', level=0)
    doc.add_paragraph('Audience: Developers, solution architects, technical support, QA')
    doc.add_paragraph(f'Date: {date.today().isoformat()}')

    doc.add_heading('1. Module Overview', level=1)
    doc.add_paragraph('Module: tender_management (Odoo 16)')
    doc.add_paragraph('Primary responsibility: tender publication, bid submission, portal tracking, workflow control, document access, and website integration.')

    doc.add_heading('2. Architecture Components', level=1)
    doc.add_paragraph('Controllers')
    doc.add_paragraph('- bid_submission.py: website listing/detail routes, submit_bid, timer API, document downloads')
    doc.add_paragraph('- vendor_registration.py: registration form handling and partner creation')
    doc.add_paragraph('- tender_bid_portal.py: portal bid list/detail and bid document downloads')
    doc.add_paragraph('Models')
    doc.add_paragraph('- tender.tender: state machine, publication flags, draft-lock behavior, cron close logic')
    doc.add_paragraph('- tender.bid: bid workflow, evaluation/adjudication actions, draft-lock behavior')
    doc.add_paragraph('- tender.document: linked document storage with parent tender lock rules')
    doc.add_paragraph('Views/Templates')
    doc.add_paragraph('- Backend forms and trees: tender_tender.xml, tender_bid.xml')
    doc.add_paragraph('- Website templates: tender_tender_template.xml, tender_information_template.xml')
    doc.add_paragraph('- Portal templates: tender_bid_portal_views.xml')
    doc.add_paragraph('Assets')
    doc.add_paragraph('- static/src/js/website_tender.js (timer JSON-RPC + archive dropdown)')
    doc.add_paragraph('- static/src/css/tender_information.css (website presentation and responsive behavior)')

    doc.add_heading('3. End-to-End Technical Flows', level=1)
    doc.add_paragraph('Tender publication flow: draft -> submit -> verify -> approve -> open')
    doc.add_paragraph('Bid flow: draft -> submitted -> compliance/evaluation/adjudication -> done/cancel')
    doc.add_paragraph('Timer flow: hidden tender token in templates -> /tender/timer JSON route -> client countdown update')
    doc.add_paragraph('Portal flow: auth=user routes with ownership checks via _can_access_bid()')

    doc.add_heading('4. Security and Controls', level=1)
    doc.add_paragraph('Route Controls')
    doc.add_paragraph('- submit_bid: auth user, POST only, CSRF enabled')
    doc.add_paragraph('- vendor_registration_submit: POST with CSRF and rate limiting')
    doc.add_paragraph('- portal bid routes: auth user + owner/internal authorization checks')
    doc.add_paragraph('- download routes: per-record entitlement checks before file streaming')
    doc.add_paragraph('Validation Controls')
    doc.add_paragraph('- Upload extensions restricted to PDF, DOC, DOCX')
    doc.add_paragraph('- Max file size 10MB per file')
    doc.add_paragraph('- Mandatory upload requirement on bid submit')
    doc.add_paragraph('- Tender state and publication checks before allowing submission')
    doc.add_paragraph('Timer Hardening')
    doc.add_paragraph('- HMAC token generation and compare_digest validation against database.secret')

    doc.add_heading('5. Data Integrity and Locking', level=1)
    doc.add_paragraph('- tender.tender.write blocks non-draft edits except controlled workflow/system fields')
    doc.add_paragraph('- tender.bid.write blocks non-draft edits except controlled fields')
    doc.add_paragraph('- tender.document create/write/unlink blocked when linked tender is non-draft')
    doc.add_paragraph('- fields_get readonly metadata applied for non-draft records in tender and bid models')

    doc.add_heading('6. Configuration Keys', level=1)
    doc.add_paragraph('- tender_management.folder_id')
    doc.add_paragraph('- tender_management.rate_limit.submit_bid.max_requests')
    doc.add_paragraph('- tender_management.rate_limit.submit_bid.window_seconds')
    doc.add_paragraph('- tender_management.rate_limit.vendor_registration.max_requests')
    doc.add_paragraph('- tender_management.rate_limit.vendor_registration.window_seconds')

    doc.add_heading('7. Troubleshooting (Technical)', level=1)
    doc.add_paragraph('Timer issues')
    doc.add_paragraph('- Verify tender_timer_token exists in rendered HTML and JSON request payload includes token.')
    doc.add_paragraph('Unauthorized downloads')
    doc.add_paragraph('- Validate tender publication status and bid ownership checks.')
    doc.add_paragraph('Submission blocked')
    doc.add_paragraph('- Check blocked states, allow_bidding flag, auth context, and upload validation failures.')
    doc.add_paragraph('Rate limit false positives')
    doc.add_paragraph('- Review IP-based bucket values and deployment topology.')

    doc.add_heading('8. Test Matrix', level=1)
    doc.add_paragraph('- Positive and negative tests for each route and state transition')
    doc.add_paragraph('- Access-control tests by persona: public, portal, internal')
    doc.add_paragraph('- Upload format/size negative tests')
    doc.add_paragraph('- Timer token tampering tests (missing token / invalid token)')

    add_flow_images(doc, '9. Graphical BPMN Process Flows')

    doc.add_heading('10. File Reference Summary', level=1)
    doc.add_paragraph('See module files under tender_management/controller, tender_management/models, tender_management/views, tender_management/security, and tender_management/static/src for implementation details.')

    doc.save(TECH_OUT)


if __name__ == '__main__':
    build_user_guide()
    build_technical_guide()
    build_user_html()
    build_technical_html()
    print(USER_OUT)
    print(TECH_OUT)
    print(USER_HTML_OUT)
    print(TECH_HTML_OUT)
